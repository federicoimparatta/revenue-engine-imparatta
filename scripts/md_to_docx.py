#!/usr/bin/env python3
"""Minimal Markdown -> .docx converter for the revenue-engine skill.

Handles: # / ## / ### headings, - / * bullets, | pipe | tables, --- rules,
**bold** inline, blank-line-separated paragraphs (internal single newlines become
line breaks, e.g. a letter signature). No external deps beyond python-docx.

Usage: python3 md_to_docx.py <input.md> <output.docx>
"""
import re
import sys

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_inline(paragraph, text):
    """Add text to a paragraph, rendering **bold** segments and \n as line breaks."""
    for line_idx, line in enumerate(text.split("\n")):
        if line_idx:
            paragraph.add_run().add_break()
        for seg_idx, seg in enumerate(re.split(r"(\*\*.+?\*\*)", line)):
            if not seg:
                continue
            if seg.startswith("**") and seg.endswith("**"):
                run = paragraph.add_run(seg[2:-2])
                run.bold = True
            else:
                paragraph.add_run(seg)


def main():
    src, dst = sys.argv[1], sys.argv[2]
    with open(src, "r", encoding="utf-8") as fh:
        lines = fh.read().split("\n")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    para_buf = []

    def flush_para():
        if para_buf:
            text = "\n".join(para_buf).strip("\n")
            if text.strip():
                add_inline(doc.add_paragraph(), text)
            para_buf.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # table block: consecutive lines starting with |
        if stripped.startswith("|"):
            flush_para()
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = []
            for r in block:
                if re.match(r"^\|[\s:\-|]+\|?$", r):  # separator row
                    continue
                cells = [c.strip() for c in r.strip("|").split("|")]
                rows.append(cells)
            if rows:
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=ncols)
                table.style = "Light Grid Accent 1"
                for ridx, r in enumerate(rows):
                    cells = table.add_row().cells
                    for c in range(ncols):
                        val = r[c] if c < len(r) else ""
                        p = cells[c].paragraphs[0]
                        add_inline(p, val)
                        if ridx == 0:
                            for run in p.runs:
                                run.bold = True
            continue

        if stripped.startswith("### "):
            flush_para(); doc.add_heading(stripped[4:], level=3); i += 1; continue
        if stripped.startswith("## "):
            flush_para(); doc.add_heading(stripped[3:], level=2); i += 1; continue
        if stripped.startswith("# "):
            flush_para(); doc.add_heading(stripped[2:], level=1); i += 1; continue
        if stripped in ("---", "***", "___"):
            flush_para(); i += 1; continue
        if re.match(r"^[-*] ", stripped):
            flush_para()
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, stripped[2:])
            i += 1; continue
        if stripped == "":
            flush_para(); i += 1; continue

        para_buf.append(line)
        i += 1

    flush_para()
    doc.save(dst)
    print(f"wrote {dst}")


if __name__ == "__main__":
    main()
